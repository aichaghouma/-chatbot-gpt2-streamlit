import streamlit as st
import torch
import re
import io
import random
import speech_recognition as sr
from gtts import gTTS
from docx import Document as DocxDocument
from transformers import GPT2LMHeadModel, GPT2Tokenizer
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS
from sklearn.metrics.pairwise import cosine_similarity
from rapidfuzz import process, fuzz

from knowledge_base import KNOWLEDGE_BASE
from deep_translator import GoogleTranslator, MyMemoryTranslator

# ============================================================
# TRADUCTION (avec secours + cache)
# ============================================================
MOTS_CLES_ERREUR_TRADUCTION = [
    "MYMEMORY WARNING", "QUERY LENGTH LIMIT", "INVALID SOURCE",
    "INVALID TARGET", "TRANSLATION UNAVAILABLE", "AVAILABLE FREE TRANSLATIONS",
]


def traduction_est_valide(texte_original, texte_traduit):
    if not texte_traduit or not texte_traduit.strip():
        return False
    if any(mot in texte_traduit.upper() for mot in MOTS_CLES_ERREUR_TRADUCTION):
        return False
    if texte_traduit.strip() == texte_original.strip():
        return False
    return True


_CODES_MYMEMORY = {"fr": "fr-FR", "en": "en-GB"}
_cache_traductions = {}


def traduire_avec_secours(texte, source, cible):
    cle = (texte, source, cible)
    if cle in _cache_traductions:
        return _cache_traductions[cle]
    try:
        resultat = GoogleTranslator(source=source, target=cible).translate(texte)
        if not traduction_est_valide(texte, resultat):
            raise ValueError("Traduction Google invalide")
        _cache_traductions[cle] = resultat
        return resultat
    except Exception:
        try:
            source_mm = _CODES_MYMEMORY.get(source, source)
            cible_mm = _CODES_MYMEMORY.get(cible, cible)
            resultat = MyMemoryTranslator(source=source_mm, target=cible_mm).translate(texte)
            if not traduction_est_valide(texte, resultat):
                raise ValueError("Traduction MyMemory invalide")
            _cache_traductions[cle] = resultat
            return resultat
        except Exception:
            raise


def traduire_en_francais(texte_anglais):
    try:
        return traduire_avec_secours(texte_anglais, "en", "fr")
    except Exception:
        return texte_anglais + "\n\n*(Traduction indisponible, réponse affichée en anglais)*"


# ============================================================
# ASR (Reconnaissance vocale) ET TTS (Synthèse vocale)
# ============================================================

def transcrire_audio(audio_bytes):
    """Transcrit un fichier audio (bytes WAV) en texte, via Google Web Speech API (gratuit)."""
    recognizer = sr.Recognizer()
    try:
        with sr.AudioFile(io.BytesIO(audio_bytes)) as source:
            audio_data = recognizer.record(source)
        try:
            return recognizer.recognize_google(audio_data, language="fr-FR")
        except sr.UnknownValueError:
            return recognizer.recognize_google(audio_data, language="en-US")
    except sr.UnknownValueError:
        return None
    except sr.RequestError:
        return None
    except Exception:
        return None


def generer_audio(texte, francais=True):
    """Convertit un texte en audio (bytes mp3) via gTTS."""
    try:
        lang = "fr" if francais else "en"
        tts = gTTS(text=texte, lang=lang)
        buffer = io.BytesIO()
        tts.write_to_fp(buffer)
        buffer.seek(0)
        return buffer.read()
    except Exception:
        return None


# ============================================================
# GENERATEUR DE QCM / EXAMEN (basé sur la base de connaissances)
# ============================================================

def generer_qcm(matiere, nb_questions=5):
    """Génère un QCM à partir des fiches d'une matière (ou de toute la base si matiere='Toutes')."""
    if matiere == "Toutes":
        pool = KNOWLEDGE_BASE
    else:
        pool = [d for d in KNOWLEDGE_BASE if d["subject"] == matiere]

    if len(pool) < 2:
        return []

    nb_questions = min(nb_questions, len(pool))
    fiches_questions = random.sample(pool, nb_questions)

    autre_pool = [d for d in pool if d not in fiches_questions]
    if len(autre_pool) < 3:
        autre_pool = [d for d in KNOWLEDGE_BASE if d not in fiches_questions]

    qcm = []
    for fiche in fiches_questions:
        distracteurs_dispo = [d for d in autre_pool if d["title"] != fiche["title"]]
        nb_distracteurs = min(3, len(distracteurs_dispo))
        distracteurs = random.sample(distracteurs_dispo, nb_distracteurs)

        options = [fiche["content"]] + [d["content"] for d in distracteurs]
        random.shuffle(options)
        bonne_reponse_idx = options.index(fiche["content"])

        qcm.append({
            "question": f"Que signifie / qu'est-ce que : {fiche['title']} ?",
            "options": options,
            "bonne_reponse_idx": bonne_reponse_idx,
            "sujet": fiche["subject"],
        })
    return qcm


# ============================================================
# DETECTION DE LANGUE ET TRADUCTION
# ============================================================

MOTS_FRANCAIS = {
    "quelle", "quel", "quels", "quelles", "qu'est-ce", "qu est ce",
    "comment", "pourquoi", "combien", "où", "ou est", "qui est",
    "est-ce", "peux-tu", "peux tu", "explique", "explique-moi",
    "définis", "définir", "c'est quoi", "c est quoi", "quelles sont",
    "quels sont", "je veux savoir", "dis-moi", "dis moi",
}


def est_francais(question):
    """Détecte si la question est posée en français (heuristique simple)."""
    q = question.lower()
    if any(c in q for c in "éèêàùçôî"):
        return True
    if any(mot in q for mot in MOTS_FRANCAIS):
        return True
    return False


# ============================================================
# TRADUCTEUR (fonctionnalité dédiée, prioritaire)
# ============================================================

MOTS_TRADUCTION = ["traduire", "traduis", "traduction de", "translate"]

MOTS_LANGUE_CIBLE = {
    "fr": ["en français", "en francais", "vers le français", "vers le francais", "to french", "into french", "in french"],
    "en": ["en anglais", "vers l'anglais", "vers langlais", "to english", "into english", "in english"],
    "ar": ["en arabe", "vers l'arabe", "vers larabe", "to arabic", "into arabic", "in arabic"],
}
NOMS_LANGUE = {"fr": "français", "en": "anglais", "ar": "arabe"}


def detecter_langue_texte(texte):
    """Détecte la langue probable d'un texte court (français, anglais ou arabe)."""
    if re.search(r'[\u0600-\u06FF]', texte):
        return "ar"
    if est_francais(texte):
        return "fr"
    return "en"


def detecter_demande_traduction(question):
    """Détecte une demande de traduction et extrait le texte + la langue cible."""
    q = question.lower()
    if not any(m in q for m in MOTS_TRADUCTION):
        return None

    cible = None
    for code, variantes in MOTS_LANGUE_CIBLE.items():
        if any(v in q for v in variantes):
            cible = code
            break

    if ":" in question:
        texte = question.split(":", 1)[1].strip()
    else:
        texte = question
        toutes_variantes_langue = [v for variantes in MOTS_LANGUE_CIBLE.values() for v in variantes]
        a_retirer = MOTS_TRADUCTION + toutes_variantes_langue + ["cette phrase", "ce texte", "this sentence", "?"]
        for m in a_retirer:
            texte = re.sub(re.escape(m), "", texte, flags=re.IGNORECASE)
        texte = texte.strip(" :.-")

    if not texte:
        return None
    return texte, cible


def executer_traduction(texte, cible):
    """Traduit un texte donné vers la langue cible (déduite automatiquement si None)."""
    try:
        source = detecter_langue_texte(texte)
        if cible is None:
            cible = "en" if source == "fr" else "fr"
        if cible == source:
            source = "auto"
        traduction = traduire_avec_secours(texte, source, cible)
        langue_nom = NOMS_LANGUE.get(cible, cible)
        affichage = f'"{texte}" → **{traduction}** ({langue_nom})'
        return affichage, traduction, (cible == "fr")
    except Exception:
        msg = "Désolé, la traduction n'a pas pu être effectuée (problème de connexion)."
        return msg, msg, True


# ============================================================
# DETECTION DE CONVERSION DE DEVISES (hors de portée, réponse honnête)
# ============================================================

MOTS_DEVISES = [
    "dinar", "euro", "dollar", "livre sterling", "franc suisse", "yen",
    "usd", "eur", "gbp", "tnd", "cad", "taux de change", "exchange rate",
    "combien vaut", "convertir en", "conversion de devise",
]


def est_demande_devise(question):
    """Détecte une demande de conversion de devises (impossible à répondre de façon fiable et à jour)."""
    q = question.lower()
    nb_mots_devise = sum(1 for m in MOTS_DEVISES if m in q)
    return nb_mots_devise >= 2


# ============================================================
# CALCULATEUR ARITHMÉTIQUE (priorité absolue avant l'IA)
# ============================================================

def calculer_expression(question):
    """Détecte et calcule une expression arithmétique complète (+, -, *, /, parenthèses)."""
    q = question.replace(" ", "")
    match = re.search(r'[\d\.\+\-\*x×/\(\)]{3,}', q)
    if not match:
        return None

    expr = match.group().rstrip("+-*/x×.")
    if not expr:
        return None

    expr_calc = expr.replace("x", "*").replace("×", "*")
    if not re.fullmatch(r'[\d\.\+\-\*/\(\)]+', expr_calc):
        return None
    if expr_calc.count("(") != expr_calc.count(")"):
        return None

    try:
        resultat = eval(expr_calc, {"__builtins__": {}}, {})
    except ZeroDivisionError:
        return "Division by zero is undefined."
    except Exception:
        return None

    if isinstance(resultat, float) and resultat == int(resultat):
        resultat = int(resultat)

    expr_affichage = expr.replace("x", "×").replace("*", "×")
    return f"{expr_affichage} = {resultat}"


# ============================================================
# MINI BASE DE FAITS VÉRIFIÉS (capitales du monde)
# ============================================================

CAPITALES = {
    "france": "Paris", "tunisia": "Tunis", "tunisie": "Tunis",
    "morocco": "Rabat", "maroc": "Rabat", "algeria": "Algiers",
    "algerie": "Algiers", "libya": "Tripoli", "egypt": "Cairo",
    "egypte": "Cairo", "germany": "Berlin", "allemagne": "Berlin",
    "italy": "Rome", "italie": "Rome", "spain": "Madrid",
    "espagne": "Madrid", "portugal": "Lisbon", "uk": "London",
    "united kingdom": "London", "england": "London",
    "angleterre": "London", "usa": "Washington, D.C.",
    "united states": "Washington, D.C.", "etats-unis": "Washington, D.C.",
    "canada": "Ottawa", "china": "Beijing", "chine": "Beijing",
    "japan": "Tokyo", "japon": "Tokyo", "india": "New Delhi",
    "inde": "New Delhi", "brazil": "Brasilia", "bresil": "Brasilia",
    "russia": "Moscow", "russie": "Moscow", "turkey": "Ankara",
    "turquie": "Ankara", "greece": "Athens", "grece": "Athens",
    "switzerland": "Bern", "suisse": "Bern", "belgium": "Brussels",
    "belgique": "Brussels", "netherlands": "Amsterdam",
    "pays-bas": "Amsterdam", "senegal": "Dakar", "mali": "Bamako",
    "ivory coast": "Yamoussoukro", "cote d'ivoire": "Yamoussoukro",
    "saudi arabia": "Riyadh", "arabie saoudite": "Riyadh",
    "qatar": "Doha", "uae": "Abu Dhabi",
    "emirats arabes unis": "Abu Dhabi",
}


def chercher_capitale(question):
    """Vérifie si la question porte sur une capitale connue et fiable."""
    q = question.lower()
    if "capital" not in q:
        return None
    for pays, capitale in CAPITALES.items():
        if pays in q:
            return f"The capital of {pays.title()} is {capitale}."
    return None


# ============================================================
# MOTEUR RAG (recherche dans la base de connaissances multi-matières)
# ============================================================

SEUIL_SIMILARITE = 0.15
SEUIL_DOC_BRUT = 0.30  # seuil plus strict pour la recherche en français brut (non traduit)


@st.cache_resource
def construire_index_rag():
    """Prépare le moteur de recherche (TF-IDF) sur la base de connaissances."""
    textes = [f"{doc['title']} {doc['content']}" for doc in KNOWLEDGE_BASE]
    stop_words_etendus = list(ENGLISH_STOP_WORDS) + [
        "define", "defined", "defining", "definition",
        "explain", "explains", "explained", "explaining",
        "describe", "describes", "described", "describing",
        "tell", "tells", "telling",
    ]
    vectorizer = TfidfVectorizer(stop_words=stop_words_etendus)
    matrix = vectorizer.fit_transform(textes)
    return vectorizer, matrix


@st.cache_resource
def construire_vocabulaire():
    """Extrait tous les mots uniques du titre+contenu de la base, pour la correction orthographique."""
    mots = set()
    for doc in KNOWLEDGE_BASE:
        texte = f"{doc['title']} {doc['content']}".lower()
        mots.update(re.findall(r"[a-zA-Zàâäéèêëïîôöùûüç]+", texte))
    return list(mots)


def corriger_question(question, vocabulaire_connu, seuil=80):
    """Corrige les mots mal orthographiés proches d'un mot connu de la base."""
    mots = question.split()
    mots_corriges = []
    for mot in mots:
        if len(mot) < 4:
            mots_corriges.append(mot)
            continue
        match = process.extractOne(mot.lower(), vocabulaire_connu, scorer=fuzz.ratio, score_cutoff=seuil)
        mots_corriges.append(match[0] if match else mot)
    return " ".join(mots_corriges)


EXPANSIONS_SYNONYMES = {
    "ml": "machine learning", "ai": "artificial intelligence", "dl": "deep learning",
    "db": "database", "os": "operating system", "ui": "user interface",
    "ux": "user experience", "api": "application programming interface",
    "nn": "neural network", "nlp": "natural language processing",
    "cs": "computer science", "iot": "internet of things",
}


def etendre_synonymes(texte):
    """Remplace les acronymes courants par leur forme complète pour améliorer le matching RAG."""
    mots = texte.split()
    mots_etendus = []
    for m in mots:
        m_propre = re.sub(r"[^a-zA-Z]", "", m).lower()
        mots_etendus.append(EXPANSIONS_SYNONYMES.get(m_propre, m))
    return " ".join(mots_etendus)


def chercher_dans_base(question, vectorizer, matrix):
    """Trouve la fiche la plus pertinente pour la question, si elle existe."""
    question_etendue = etendre_synonymes(question)
    q_vec = vectorizer.transform([question_etendue])
    scores = cosine_similarity(q_vec, matrix)[0]
    idx = scores.argmax()
    if scores[idx] >= SEUIL_SIMILARITE:
        return KNOWLEDGE_BASE[idx], scores[idx]
    return None, scores[idx]


def generer_reponse_avec_contexte(model, tokenizer, question, contexte, device, max_length=100, temperature=0.4):
    """Génère une réponse en s'appuyant sur un passage de la base de connaissances (RAG)."""
    model.eval()
    prompt = f"Context: {contexte}\nQuestion: {question}\nAnswer:"
    inputs = tokenizer(prompt, return_tensors="pt").to(device)

    with torch.no_grad():
        outputs = model.generate(
            **inputs,
            max_length=inputs["input_ids"].shape[1] + max_length,
            temperature=temperature,
            do_sample=True,
            top_p=0.85,
            repetition_penalty=1.15,
            no_repeat_ngram_size=3,
            pad_token_id=tokenizer.eos_token_id
        )

    reponse = tokenizer.decode(outputs[0], skip_special_tokens=True)
    reponse = reponse.split("Answer:")[-1].strip()
    phrases = re.split(r'(?<=[.!?])\s+', reponse)
    reponse_courte = " ".join(phrases[:2]).strip()
    return reponse_courte if reponse_courte else reponse


# ============================================================
# CONFIGURATION DE LA PAGE
# ============================================================
st.set_page_config(page_title="Chatbot ENET'Com - GPT-2 Fine-tuné", page_icon="🤖")

# ============================================================
# CHARGEMENT DU MODÈLE (mis en cache pour la rapidité)
# ============================================================

MODEL_NAME = "Aicha83/chatbot-gpt2-finetuned"


@st.cache_resource
def charger_modele():
    device = torch.device("cpu")
    tokenizer = GPT2Tokenizer.from_pretrained(MODEL_NAME)
    model = GPT2LMHeadModel.from_pretrained(MODEL_NAME).to(device)
    model.eval()
    return model, tokenizer, device


# ============================================================
# FONCTION DE GÉNÉRATION DE RÉPONSE (gardée disponible, non utilisée
# dans la logique de décision principale — voir corrections apportées)
# ============================================================

def generer_reponse(model, tokenizer, question, device, max_length=80, temperature=0.4):
    model.eval()
    prompt = f"Question: {question}\nAnswer:"
    inputs = tokenizer(prompt, return_tensors="pt").to(device)

    with torch.no_grad():
        outputs = model.generate(
            **inputs,
            max_length=max_length,
            temperature=temperature,
            do_sample=True,
            top_p=0.85,
            repetition_penalty=1.15,
            no_repeat_ngram_size=3,
            pad_token_id=tokenizer.eos_token_id
        )

    reponse = tokenizer.decode(outputs[0], skip_special_tokens=True)
    reponse = reponse.split("Answer:")[-1].strip()
    phrases = re.split(r'(?<=[.!?])\s+', reponse)
    reponse_courte = " ".join(phrases[:2]).strip()
    return reponse_courte if reponse_courte else reponse


# ============================================================
# INTERFACE STREAMLIT
# ============================================================

st.title("🤖 Chatbot GPT-2 Fine-tuné + RAG")
st.caption("Projet de stage — ESSE Lab, ENET'Com Sfax — GPT-2 base fine-tuné + recherche documentaire (RAG)")

with st.expander("ℹ️ À propos de ce chatbot"):
    st.markdown(f"""
    Ce chatbot utilise **GPT-2 base**, un modèle pré-entraîné par OpenAI (124M paramètres),
    **fine-tuné** sur le dataset Databricks Dolly 15k, combiné à un système **RAG**
    (Retrieval-Augmented Generation).

    **Comment ça marche :**
    1. La question est comparée automatiquement à une base de **{len(KNOWLEDGE_BASE)} fiches**
       (maths, physique, géographie, civilisation, électronique, informatique, cybersécurité, chimie)
    2. Si une fiche pertinente est trouvée → la réponse **vérifiée** de cette fiche est affichée directement (RAG extractif)
    3. Les questions de capitales sont répondues directement depuis une base fiable
    4. Sinon → réponse honnête indiquant que le sujet n'est pas couvert

    **Pourquoi un RAG "extractif" plutôt que "génératif" ?**
    Des tests ont montré que GPT-2 base, n'étant pas entraîné à l'instruction-following,
    n'arrive pas à reformuler fidèlement un contexte fourni — il continue à halluciner
    même avec la bonne information sous les yeux. Afficher directement le contenu vérifié
    de la fiche garantit donc une réponse fiable pour les sujets couverts par la base.

    **Caractéristiques techniques :**
    - Modèle de base : GPT-2 (124M paramètres)
    - Fine-tuning : 8 époques sur ~8500 exemples (Dolly 15k)
    - Recherche : TF-IDF + similarité cosinus sur la base de connaissances
    - Correction orthographique légère avant recherche (rapidfuzz)
    - Génération : anti-répétition activée (repetition_penalty, no_repeat_ngram_size)

    ⚠️ **Limites connues** : la base de connaissances est volontairement limitée (démonstration
    de projet de stage). Pour les sujets non couverts, le chatbot l'indique honnêtement
    plutôt que de générer une réponse non fiable.
    """)

# ============================================================
# SECTION QCM / EXAMEN
# ============================================================

with st.expander("📝 Générer un QCM / Examen"):
    st.markdown("Génère un questionnaire à choix multiples à partir de la base de connaissances, pour réviser une matière.")

    matieres_disponibles = ["Toutes"] + sorted(set(d["subject"] for d in KNOWLEDGE_BASE))
    col1, col2 = st.columns(2)
    with col1:
        matiere_choisie = st.selectbox("Matière", matieres_disponibles)
    with col2:
        nb_questions_choisi = st.slider("Nombre de questions", min_value=3, max_value=15, value=5)

    if st.button("🎲 Générer le QCM"):
        st.session_state.qcm_actuel = generer_qcm(matiere_choisie, nb_questions_choisi)
        st.session_state.qcm_reponses = {}
        st.session_state.qcm_corrige = False

    if st.session_state.get("qcm_actuel"):
        qcm = st.session_state.qcm_actuel
        st.markdown(f"**{len(qcm)} questions — matière : {matiere_choisie}**")

        for i, q in enumerate(qcm):
            st.markdown(f"**{i+1}. {q['question']}** *({q['sujet']})*")
            choix = st.radio(
                f"question_{i}",
                options=list(range(len(q["options"]))),
                format_func=lambda idx, opts=q["options"]: opts[idx],
                key=f"qcm_radio_{i}",
                label_visibility="collapsed",
            )
            st.session_state.qcm_reponses[i] = choix
            st.markdown("---")

        if st.button("✅ Corriger le QCM"):
            st.session_state.qcm_corrige = True

        if st.session_state.get("qcm_corrige"):
            score = sum(
                1 for i, q in enumerate(qcm)
                if st.session_state.qcm_reponses.get(i) == q["bonne_reponse_idx"]
            )
            st.success(f"Score : {score} / {len(qcm)}")
            for i, q in enumerate(qcm):
                bonne = q["bonne_reponse_idx"]
                donnee = st.session_state.qcm_reponses.get(i)
                if donnee == bonne:
                    st.markdown(f"✅ Question {i+1} : correcte")
                else:
                    st.markdown(f"❌ Question {i+1} : incorrecte — bonne réponse : *{q['options'][bonne]}*")


# Charger le modèle et l'index RAG (une seule fois, mis en cache)
try:
    model, tokenizer, device = charger_modele()
    vectorizer, matrix = construire_index_rag()
    vocabulaire_connu = construire_vocabulaire()
    modele_charge = True
except Exception as e:
    modele_charge = False
    st.error(f"Erreur lors du chargement du modèle : {e}")
    st.info(f"Vérifie que le modèle `{MODEL_NAME}` est bien public sur Hugging Face et accessible.")

if modele_charge:
    if "messages" not in st.session_state:
        st.session_state.messages = []
    if "lire_audio" not in st.session_state:
        st.session_state.lire_audio = False

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
            if message["role"] == "assistant" and "audio" in message and message["audio"]:
                st.audio(message["audio"], format="audio/mp3")

    st.session_state.lire_audio = st.checkbox(
        "🔊 Lire les réponses à voix haute", value=st.session_state.lire_audio
    )

    with st.expander("🎤 Poser la question à l'oral"):
        audio_input = st.audio_input("Enregistre ta question")

    question_vocale = None
    if audio_input is not None:
        with st.spinner("Transcription de l'audio..."):
            question_vocale = transcrire_audio(audio_input.getvalue())
        if question_vocale is None:
            st.warning("Impossible de comprendre l'audio. Réessaie ou pose ta question par écrit.")

    question_texte = st.chat_input("Pose ta question ici (en anglais)...")
    question = question_texte or question_vocale

    if question:
        with st.chat_message("user"):
            st.markdown(question)
        st.session_state.messages.append({"role": "user", "content": question})

        with st.chat_message("assistant"):
            with st.spinner("Génération de la réponse..."):
                francais = est_francais(question)

                question_recherche = question
                if francais:
                    try:
                        question_recherche = traduire_avec_secours(question, "fr", "en")
                    except Exception:
                        question_recherche = question

                reponse_calcul = calculer_expression(question)
                demande_trad = detecter_demande_traduction(question)
                demande_devise = est_demande_devise(question)
                texte_audio = None
                audio_est_francais = francais

                if demande_devise:
                    if francais:
                        reponse = ("Je ne peux pas faire de conversion de devises fiable, car les taux de change "
                                   "changent en temps réel et je n'ai pas accès à des données à jour. "
                                   "Utilise un convertisseur en ligne (comme Google, XE.com, ou ton application bancaire) "
                                   "pour un taux exact et actuel.")
                    else:
                        reponse = ("I can't provide reliable currency conversion, since exchange rates change "
                                   "in real time and I don't have access to live data. "
                                   "Please use an online converter (like Google, XE.com, or your banking app) "
                                   "for an exact, current rate.")
                    badge = "⚠️ Hors de portée (données en temps réel non disponibles)"

                elif reponse_calcul:
                    reponse = reponse_calcul
                    badge = "🧮 Calcul exact (Python)"

                elif demande_trad:
                    texte, cible = demande_trad
                    reponse, texte_audio, audio_est_francais = executer_traduction(texte, cible)
                    badge = "🌐 Traduction"

                elif chercher_capitale(question_recherche):
                    reponse = chercher_capitale(question_recherche)
                    if francais:
                        reponse = traduire_en_francais(reponse)
                    badge = "✅ Réponse vérifiée (base de capitales)"

                else:
                    # Si la traduction FR->EN a échoué, question_recherche est restée en français.
                    traduction_a_echoue = francais and (question_recherche == question)

                    question_recherche_corrigee = corriger_question(question_recherche, vocabulaire_connu)
                    doc_trad, score_trad = chercher_dans_base(question_recherche_corrigee, vectorizer, matrix)

                    if traduction_a_echoue and score_trad < SEUIL_DOC_BRUT:
                        doc_trad, score_trad = None, 0

                    if francais:
                        doc_brut, score_brut = chercher_dans_base(question, vectorizer, matrix)
                        if doc_brut and doc_brut["subject"] in ("French", "English"):
                            doc_brut, score_brut = None, 0
                        if doc_brut and score_brut >= SEUIL_DOC_BRUT and score_brut > score_trad + 0.1:
                            doc_trouve, score = doc_brut, score_brut
                        else:
                            doc_trouve, score = doc_trad, score_trad
                    else:
                        doc_trouve, score = doc_trad, score_trad

                    if doc_trouve:
                        reponse = doc_trouve["content"]
                        if francais:
                            reponse = traduire_en_francais(reponse)
                        badge = f"📚 Réponse vérifiée : *{doc_trouve['title']}* ({doc_trouve['subject']}) — RAG"
                    else:
                        if francais:
                            reponse = ("Je n'ai pas d'information vérifiée sur ce sujet dans ma base de "
                                       "connaissances. Essaie de reformuler ta question, ou pose une "
                                       "question sur un des sujets couverts (maths, physique, géographie, "
                                       "civilisation, électronique, informatique, cybersécurité, chimie).")
                        else:
                            reponse = ("I don't have verified information on this topic in my knowledge "
                                       "base. Try rephrasing your question, or ask about a covered subject "
                                       "(math, physics, geography, civics, electronics, computer science, "
                                       "cybersecurity, chemistry).")
                        badge = "❓ Sujet non couvert par la base de connaissances"

                st.caption(badge)
            st.markdown(reponse)

            audio_reponse = None
            if st.session_state.lire_audio:
                if texte_audio is None:
                    texte_audio = reponse
                with st.spinner("Génération de l'audio..."):
                    audio_reponse = generer_audio(texte_audio, francais=audio_est_francais)
                if audio_reponse:
                    st.audio(audio_reponse, format="audio/mp3")
                else:
                    st.caption("⚠️ Synthèse vocale indisponible pour cette réponse.")

        st.session_state.messages.append({"role": "assistant", "content": reponse, "audio": audio_reponse})

    col_export, col_reset = st.columns(2)

    with col_export:
        if st.session_state.messages:
            def generer_docx_conversation(messages):
                doc = DocxDocument()
                doc.add_heading("Conversation — Chatbot GPT-2 + RAG", level=1)
                for msg in messages:
                    role = "Vous" if msg["role"] == "user" else "Chatbot"
                    p = doc.add_paragraph()
                    p.add_run(f"{role} : ").bold = True
                    p.add_run(msg["content"])
                buffer = io.BytesIO()
                doc.save(buffer)
                buffer.seek(0)
                return buffer

            docx_buffer = generer_docx_conversation(st.session_state.messages)
            st.download_button(
                "📄 Exporter la conversation (Word)",
                data=docx_buffer,
                file_name="conversation_chatbot.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    with col_reset:
        if st.session_state.messages:
            if st.button("🗑️ Réinitialiser la conversation"):
                st.session_state.messages = []
                st.rerun()
